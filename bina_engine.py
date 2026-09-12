"""
BinaLeadPro — Multi-Source Extraction Engine (v3.0)
===================================================
Supports bina.az, yeniemlak.az, and tap.az.
Handles Rayon and Sub-location hierarchy.
Per-platform target counts, deduplication, and export to styled Excel, CSV, and JSON.
"""

import sys
import os
import time
import random
import json
import re
from typing import Dict, List, Optional, Any, Callable
import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Console encoding fix
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8', line_buffering=True)
        sys.stderr.reconfigure(encoding='utf-8', line_buffering=True)
    except Exception:
        pass

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36',
    'Accept': '*/*',
    'Accept-Language': 'az,en-US;q=0.9,en;q=0.8,ru;q=0.7',
    'Content-Type': 'application/json',
    'Referer': 'https://bina.az/',
    'Origin': 'https://bina.az'
}

GQL_URL = 'https://bina.az/graphql'

SEARCH_QUERY = """
query SearchItems($first: Int, $filter: ItemFilter, $sort: ItemConnectionSort!, $cursor: String) {
  itemsConnection(first: $first, after: $cursor, filter: $filter, sort: $sort) {
    totalCount
    pageInfo {
      hasNextPage
      endCursor
    }
    edges {
      node {
        id
        path
        isBusiness
        rooms
        floor
        floors
        updatedAt
        hasBillOfSale
        hasMortgage
        hasRepair
        area {
          units
          value
        }
        price {
          currency
          total
        }
        company {
          id
          name
          targetType
        }
        location {
          id
          name
          fullName
        }
      }
    }
  }
}
"""

ITEM_DETAIL_QUERY = """
query Item($id: ID!) {
  item(id: $id) {
    id
    address
    buildingTypeName
    contactTypeName
    contactName
    updatedAt
    rooms
    floor
    floors
    hasBillOfSale
    hasMortgage
    hasRepair
    description
    area {
      units
      value
    }
    price {
      currency
      total
      perAre
    }
    metaTags {
      name
      content
    }
    company {
      id
      name
      targetType
    }
    location {
      fullName
      path
    }
    business {
      ... on Agency {
        id
        name
      }
      ... on Residence {
        id
        name
      }
    }
  }
}
"""

AGENT_KEYWORDS = [
    'development', 'mtk', 'əmlak', 'emlak', 'group', 'qrup', 'construction',
    'tikinti', 'agent', 'agency', 'agentlik', 'realtor', 'rieltor', 'makler',
    'şirkət', 'sirket', 'kompleks', 'holding', 'brok', 'invest', 'daşınmaz',
    'dasinmaz', 'ev alqı', 'ofis', 'bina', 'estate', 'realty', 'service',
    'xidmət', 'xidmet', 'menzil', 'mənzil', 'evlərin', 'evlerin'
]

HIERARCHY_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "bina_locations_hierarchy.json")
if os.path.exists(HIERARCHY_FILE):
    try:
        with open(HIERARCHY_FILE, "r", encoding="utf-8") as f:
            HIERARCHY_MAP = json.load(f)
    except Exception:
        HIERARCHY_MAP = {}
else:
    HIERARCHY_MAP = {}

CATEGORIES_MAP = {
    "Bütün Mənzillər": 1,
    "Yeni tikili": 2,
    "Köhnə tikili": 3,
    "Həyət evi / Villa": 5,
    "Ofis": 7,
    "Qaraj": 8,
    "Torpaq": 9,
    "Obyekt": 10
}

TAP_GET_AD_QUERY = """
query GetAd($legacyId: ID!, $source: SourceEnum!) {
  adDetails(legacyId: $legacyId, source: $source) {
    id
    title
    price
    path
    contact {
      name
    }
    user {
      id
      publishedAdsCount
    }
    azProperties {
      name
      value
    }
  }
}
"""

TAP_CREATE_CALL_MUTATION = """
mutation CreateCall($adId: ID, $shopId: ID, $source: SourceEnum!) {
  createCall(adId: $adId, shopId: $shopId, source: $source) {
    entity
  }
}
"""

YENIEMLAK_RAYON_MAP = {
    "Binəqədi r.": "1",
    "Nizami r.": "2",
    "Nərimanov r.": "3",
    "Nəsimi r.": "4",
    "Qaradağ r.": "5",
    "Sabunçu r.": "6",
    "Suraxanı r.": "7",
    "Səbail r.": "8",
    "Xətai r.": "9",
    "Xəzər r.": "10",
    "Yasamal r.": "11",
    "Pirallahı r.": "36"
}


class BinaScraperEngine:
    def __init__(self, callback: Optional[Callable[[str, Any], None]] = None):
        self.callback = callback or (lambda event, data: None)
        self.session = requests.Session()
        self.session.headers.update(HEADERS)
        self.is_cancelled = False
        self.stats = {
            "scanned": 0,
            "accepted": 0,
            "skipped": 0,
            "duplicates": 0
        }
        self.phone_seen = set()

    def cancel(self):
        self.is_cancelled = True
        self.callback("log", "[!] Axtarış istifadəçi tərəfindən dayandırılır...")

    def _sleep_jitter(self, min_s: float = 0.8, max_s: float = 1.3):
        time.sleep(random.uniform(min_s, max_s))

    def fetch_search_page(self, filter_obj: Dict[str, Any], cursor: Optional[str] = None, batch_size: int = 24) -> Dict[str, Any]:
        payload = {
            'operationName': 'SearchItems',
            'query': SEARCH_QUERY,
            'variables': {
                'first': batch_size,
                'filter': filter_obj,
                'sort': 'BUMPED_AT_DESC',
                'cursor': cursor
            }
        }
        try:
            resp = self.session.post(GQL_URL, json=payload, timeout=15)
            if resp.status_code == 200:
                data = resp.json()
                return data.get('data', {}).get('itemsConnection', {})
            return {}
        except Exception as e:
            self.callback("log", f"[!] Sorğu xətası: {e}")
            return {}

    def fetch_item_detail(self, item_id: str) -> Optional[Dict[str, Any]]:
        payload = {
            'operationName': 'Item',
            'query': ITEM_DETAIL_QUERY,
            'variables': {'id': str(item_id)}
        }
        try:
            resp = self.session.post(GQL_URL, json=payload, timeout=15)
            if resp.status_code == 200:
                return resp.json().get('data', {}).get('item')
            return None
        except Exception:
            return None

    def fetch_unmasked_phone(self, item_id: str) -> Optional[str]:
        url = f"https://bina.az/items/{item_id}/phones?react=true"
        try:
            resp = self.session.get(url, timeout=10)
            if resp.status_code == 200:
                phones = resp.json().get('phones', [])
                if phones and isinstance(phones, list) and len(phones) > 0:
                    return phones[0]
            return None
        except Exception:
            return None

    def is_agent_by_name(self, name: str) -> bool:
        if not name:
            return False
        name_lower = name.lower()
        for kw in AGENT_KEYWORDS:
            if kw in name_lower:
                return True
        return False

    def scrape_yeniemlak(self, target_type: str, max_items: int, leads: List[Dict[str, Any]],
                         region_name: str = "Bütün Bakı", sublocation_name: Optional[str] = None,
                         listing_type: str = "Alqı-satqı",
                         has_bill_of_sale: Optional[bool] = None,
                         has_mortgage: Optional[bool] = None,
                         has_repair: Optional[bool] = None) -> int:
        self.callback("log", f"\n🌐 [yeniemlak.az] Bazadan {max_items} ədəd elan axtarılır...")
        y_headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
        start_count = len(leads)

        # elan_nov: 1 = Satış, 2 = Kirayə
        elan_nov = "2" if listing_type == "Kirayə" else "1"
        base_url = f"https://yeniemlak.az/elan/axtar?elan_nov={elan_nov}&emlak=1"

        rayon_id = YENIEMLAK_RAYON_MAP.get(region_name)
        if rayon_id:
            base_url += f"&seher[]=7&rayon[]={rayon_id}"
        else:
            base_url += "&seher[]=1"

        page = 1
        max_pages = 10
        try:
            while (len(leads) - start_count) < max_items and page <= max_pages and not self.is_cancelled:
                page_url = base_url if page == 1 else f"{base_url}&sehife={page}"
                resp = requests.get(page_url, headers=y_headers, timeout=12)
                soup = BeautifulSoup(resp.text, 'html.parser')

                links = []
                for a in soup.find_all('a', href=True):
                    href = a['href']
                    if re.search(r'/elan/[a-z0-9-]+-\d+$', href):
                        full_url = href if href.startswith('http') else 'https://yeniemlak.az' + href
                        if full_url not in links:
                            links.append(full_url)

                if not links:
                    break

                for link in links:
                    if self.is_cancelled or (len(leads) - start_count) >= max_items:
                        break

                    self.stats["scanned"] += 1
                    self.callback("stats", self.stats)
                    self._sleep_jitter(0.4, 0.8)

                    try:
                        r_item = requests.get(link, headers=y_headers, timeout=10)
                        soup_item = BeautifulSoup(r_item.text, 'html.parser')
                        item_text = soup_item.text

                        # Check special conditions if requested
                        has_k = "Kupça" in item_text
                        has_i = "İpoteka" in item_text
                        has_t = "Təmirli" in item_text or "əla təmir" in item_text.lower()

                        if has_bill_of_sale and not has_k:
                            self.stats["skipped"] += 1
                            continue
                        if has_mortgage and not has_i:
                            self.stats["skipped"] += 1
                            continue
                        if has_repair and not has_t:
                            self.stats["skipped"] += 1
                            continue

                        img_tel = soup_item.find('img', src=re.compile(r'/tel-show/'))
                        if not img_tel:
                            self.stats["skipped"] += 1
                            continue

                        m = re.search(r'/tel-show/(\d+)', img_tel['src'])
                        if not m:
                            self.stats["skipped"] += 1
                            continue

                        raw_phone = m.group(1)
                        phone_clean = raw_phone
                        if phone_clean in self.phone_seen or len(raw_phone) < 7:
                            self.stats["duplicates"] += 1
                            self.callback("stats", self.stats)
                            continue

                        elvrn = soup_item.find('div', class_='elvrn')
                        status_text = elvrn.text.strip() if elvrn else ""

                        is_owner = "əmlak sahibi" in status_text.lower() or "sahibi" in status_text.lower()
                        if target_type == "FSBO" and not is_owner:
                            self.stats["skipped"] += 1
                            continue
                        elif target_type == "AGENT" and is_owner:
                            self.stats["skipped"] += 1
                            continue

                        self.phone_seen.add(phone_clean)
                        formatted_phone = f"({raw_phone[:3]}) {raw_phone[3:6]}-{raw_phone[6:8]}-{raw_phone[8:]}" if len(raw_phone) == 10 else raw_phone

                        ad_div = soup_item.find('div', class_='ad')
                        name_str = ad_div.text.strip() if ad_div else ("Ev Sahibi" if target_type == "FSBO" else "Vasitəçi")

                        price_tag = soup_item.find('price')
                        if price_tag and price_tag.text.strip():
                            price_val = float(re.sub(r'[^\d.]', '', price_tag.text.strip()))
                        else:
                            price_match = re.search(r'(\d+[\s\d]*)\s*AZN', soup_item.text)
                            price_val = float(price_match.group(1).replace(' ', '')) if price_match else 0

                        rooms_match = re.search(r'(\d+)\s*otaq', soup_item.text, re.IGNORECASE)
                        rooms_val = int(rooms_match.group(1)) if rooms_match else "-"

                        area_match = re.search(r'(\d+(?:\.\d+)?)\s*(?:m2|m²)', soup_item.text)
                        area_val = float(area_match.group(1)) if area_match else "-"

                        unvan_match = re.search(r'Ünvan\s*\n\s*([^\n]+)', soup_item.text)
                        loc_match = unvan_match.group(1).strip() if unvan_match else (region_name if region_name != "Bütün Bakı" else "Bakı")

                        # Price Trend & SQM Price
                        price_trend = "Bazar Qiyməti"
                        item_lower = item_text.lower()
                        if any(w in item_lower for w in ["endirim", "endirildi", "qiymət düşdü", "qiymeti dusdu", "ucuz"]):
                            price_trend = "🔻 Endirimli"
                        elif any(w in item_lower for w in ["təcili", "tecili", "təcili satılır"]):
                            price_trend = "🔥 Təcili Satış"

                        sqm_val = round(price_val / float(area_val)) if area_val != '-' and float(area_val) > 0 and price_val > 0 else "-"

                        record = {
                            "№": len(leads) + 1,
                            "Mənbə": "yeniemlak.az",
                            "Ad / Mülkiyyətçi": name_str,
                            "Əlaqə Nömrəsi": formatted_phone,
                            "Status": "Mülkiyyətçi" if target_type == "FSBO" else "Vasitəçi (agent)",
                            "Agentlik / Şirkət": "-" if target_type == "FSBO" else name_str,
                            "Qiymət (AZN)": price_val,
                            "m² Qiyməti (AZN)": sqm_val,
                            "Qiymət Trendi": price_trend,
                            "Otaq": rooms_val,
                            "Sahə (m²)": area_val,
                            "Yerləşmə": loc_match,
                            "Kupça": "Var" if has_k else "Yoxdur",
                            "İpoteka": "Var" if has_i else "Yoxdur",
                            "Təmir": "Var" if has_t else "Yoxdur",
                            "Elan ID": re.search(r'\d+$', link).group(0),
                            "Elan Keçidi": link,
                            "Tarix": time.strftime("%Y-%m-%d")
                        }

                        leads.append(record)
                        self.stats["accepted"] = len(leads)
                        self.callback("stats", self.stats)
                        self.callback("lead", record)
                        self.callback("log", f"  ✅ [yeniemlak.az #{len(leads) - start_count}/{max_items}] {name_str} | {formatted_phone} | {price_val} AZN")
                    except Exception:
                        continue
                    except Exception:
                        continue

                page += 1

            return len(leads) - start_count
        except Exception as e:
            self.callback("log", f"[!] yeniemlak.az xətası: {e}")
            return 0

    def scrape_tapaz(self, target_type: str, max_items: int, leads: List[Dict[str, Any]],
                     region_name: str = "Bütün Bakı", sublocation_name: Optional[str] = None,
                     category_name: str = "Bütün Mənzillər", listing_type: str = "Alqı-satqı") -> int:
        """Extracts listings from tap.az real estate with 100% unmasked phone numbers via GraphQL"""
        self.callback("log", f"\n🌐 [tap.az] Bazadan {max_items} ədəd elan axtarılır...")
        start_count = len(leads)

        tap_graphql_headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36',
            'Content-Type': 'application/json',
            'Origin': 'https://tap.az',
            'Referer': 'https://tap.az/'
        }

        cat_slug = "menziller"
        if "Villa" in category_name or "Həyət evi" in category_name:
            cat_slug = "heyet-evleri"
        elif "Ofis" in category_name or "Obyekt" in category_name:
            cat_slug = "obyektler-ve-ofisler"
        elif "Torpaq" in category_name:
            cat_slug = "torpaq-sahesi"

        keyword = ""
        if sublocation_name and sublocation_name != "Bütün rayon üzrə":
            keyword = sublocation_name.replace(" m.", "").replace(" qəs.", "").strip()
        elif region_name and region_name != "Bütün Bakı":
            keyword = region_name.replace(" r.", "").strip()

        page = 1
        max_pages = 10

        try:
            while (len(leads) - start_count) < max_items and page <= max_pages and not self.is_cancelled:
                search_url = f"https://tap.az/elanlar/dasinmaz-emlak/{cat_slug}?page={page}"
                if keyword:
                    search_url += f"&keywords={keyword}"

                resp = requests.get(search_url, headers={'User-Agent': 'Mozilla/5.0'}, timeout=12)
                soup = BeautifulSoup(resp.text, 'html.parser')

                links = [a['href'] for a in soup.find_all('a', href=True) if '/elanlar/dasinmaz-emlak/' in a['href'] and re.search(r'\d+$', a['href'])]
                unique_links = list(dict.fromkeys(links))

                if not unique_links:
                    break

                for link in unique_links:
                    if self.is_cancelled or (len(leads) - start_count) >= max_items:
                        break

                    legacy_id = re.search(r'\d+$', link).group(0)
                    full_link = link if link.startswith('http') else 'https://tap.az' + link

                    self.stats["scanned"] += 1
                    self.callback("stats", self.stats)
                    self._sleep_jitter(0.5, 0.9)

                    try:
                        # 1. Fetch ad details
                        r_ad = requests.post("https://tap.az/graphql", headers=tap_graphql_headers, json={
                            "operationName": "GetAd",
                            "query": TAP_GET_AD_QUERY,
                            "variables": {"legacyId": legacy_id, "source": "DESKTOP"}
                        }, timeout=10)

                        ad_data = r_ad.json().get("data", {}).get("adDetails")
                        if not ad_data:
                            self.stats["skipped"] += 1
                            continue

                        ad_id = ad_data.get("id")
                        title = ad_data.get("title", "")
                        price_val = float(ad_data.get("price") or 0)
                        name_str = (ad_data.get("contact") or {}).get("name", "Elan Sahibi")
                        pub_ads = (ad_data.get("user") or {}).get("publishedAdsCount", 1)

                        is_agent = (pub_ads > 2) or self.is_agent_by_name(name_str)
                        if target_type == "FSBO" and is_agent:
                            self.stats["skipped"] += 1
                            continue
                        elif target_type == "AGENT" and not is_agent:
                            self.stats["skipped"] += 1
                            continue

                        az_props = ad_data.get("azProperties") or []
                        elan_novu = ""
                        bina_tipi = ""
                        az_rooms = None
                        az_area = None
                        az_loc = None

                        for p in az_props:
                            p_name = (p.get("name") or "").strip().lower()
                            p_val = (p.get("value") or "").strip()
                            if "elan növü" in p_name:
                                elan_novu = p_val.lower()
                            elif "binanın tipi" in p_name:
                                bina_tipi = p_val
                            elif "otaq sayı" in p_name:
                                try:
                                    az_rooms = int(p_val)
                                except Exception:
                                    pass
                            elif "sahə" in p_name:
                                try:
                                    az_area = float(p_val.replace(',', '.'))
                                except Exception:
                                    pass
                            elif "yerləşmə yeri" in p_name:
                                az_loc = p_val

                        # Determine if this ad is rental or sale
                        title_lower = title.lower()
                        is_rent = ("kirayə" in elan_novu) or ("kiraye" in elan_novu) or ("kirayə" in title_lower) or ("kiraye" in title_lower)
                        is_sale = ("sat" in elan_novu) or ("satılır" in title_lower) or ("satilir" in title_lower)

                        if listing_type == "Alqı-satqı" and is_rent:
                            self.stats["skipped"] += 1
                            continue
                        elif listing_type == "Kirayə" and is_sale and not is_rent:
                            self.stats["skipped"] += 1
                            continue

                        # Filter by category if specific
                        if category_name == "Yeni tikili" and bina_tipi and "yeni" not in bina_tipi.lower():
                            self.stats["skipped"] += 1
                            continue
                        elif category_name == "Köhnə tikili" and bina_tipi and "köhnə" not in bina_tipi.lower():
                            self.stats["skipped"] += 1
                            continue

                        # 2. Call mutation to get 100% UNMASKED phone number
                        r_call = requests.post("https://tap.az/graphql", headers=tap_graphql_headers, json={
                            "operationName": "CreateCall",
                            "query": TAP_CREATE_CALL_MUTATION,
                            "variables": {"adId": ad_id, "shopId": None, "source": "DESKTOP"}
                        }, timeout=10)

                        phones_list = r_call.json().get("data", {}).get("createCall", {}).get("entity", [])
                        raw_phone = phones_list[0] if phones_list else ""
                        if not raw_phone or len(raw_phone) < 7:
                            self.stats["skipped"] += 1
                            continue

                        phone_clean = re.sub(r'[\s\(\)-]', '', raw_phone)
                        if phone_clean in self.phone_seen:
                            self.stats["duplicates"] += 1
                            self.callback("stats", self.stats)
                            continue

                        self.phone_seen.add(phone_clean)

                        rooms_m = re.search(r'(\d+)\s*-\s*otaqlı', title)
                        rooms_val = az_rooms if az_rooms is not None else (int(rooms_m.group(1)) if rooms_m else "-")

                        area_m = re.search(r'(\d+(?:\.\d+)?)\s*m²', title)
                        area_val = az_area if az_area is not None else (float(area_m.group(1)) if area_m else "-")

                        loc_display = az_loc if az_loc else (keyword if keyword else (region_name if region_name != "Bütün Bakı" else "Bakı"))

                        price_trend = "🔻 Endirimli" if "endirim" in title.lower() else ("🔥 Təcili Satış" if "təcili" in title.lower() else "Bazar Qiyməti")
                        sqm_val = round(price_val / float(area_val)) if area_val != '-' and float(area_val) > 0 and price_val > 0 else "-"

                        record = {
                            "№": len(leads) + 1,
                            "Mənbə": "tap.az",
                            "Ad / Mülkiyyətçi": name_str,
                            "Əlaqə Nömrəsi": raw_phone,
                            "Status": "Mülkiyyətçi" if target_type == "FSBO" else "Vasitəçi (agent)",
                            "Agentlik / Şirkət": "-" if target_type == "FSBO" else (name_str if self.is_agent_by_name(name_str) else "Fərdi Vasitəçi"),
                            "Qiymət (AZN)": price_val,
                            "m² Qiyməti (AZN)": sqm_val,
                            "Qiymət Trendi": price_trend,
                            "Otaq": rooms_val,
                            "Sahə (m²)": area_val,
                            "Yerləşmə": loc_display,
                            "Kupça": "-",
                            "İpoteka": "-",
                            "Təmir": "-",
                            "Elan ID": legacy_id,
                            "Elan Keçidi": full_link,
                            "Tarix": time.strftime("%Y-%m-%d")
                        }

                        leads.append(record)
                        self.stats["accepted"] = len(leads)
                        self.callback("stats", self.stats)
                        self.callback("lead", record)
                        self.callback("log", f"  ✅ [tap.az #{len(leads) - start_count}/{max_items}] {name_str} | {raw_phone} | {price_val} AZN")
                    except Exception:
                        continue

                page += 1

            return len(leads) - start_count
        except Exception as e:
            self.callback("log", f"[!] tap.az xətası: {e}")
            return 0

    def scrape_bina(self,
                    target_type: str,
                    max_items: int,
                    leads: List[Dict[str, Any]],
                    region_name: str = "Bütün Bakı",
                    sublocation_name: Optional[str] = None,
                    category_name: str = "Bütün Mənzillər",
                    listing_type: str = "Alqı-satqı",
                    room_ids: Optional[List[str]] = None,
                    price_min: Optional[float] = None,
                    price_max: Optional[float] = None,
                    area_min: Optional[float] = None,
                    area_max: Optional[float] = None,
                    has_bill_of_sale: Optional[bool] = None,
                    has_mortgage: Optional[bool] = None,
                    has_repair: Optional[bool] = None) -> int:
        start_count = len(leads)
        if max_items <= 0 or self.is_cancelled:
            return 0

        self.callback("log", f"\n🌐 [bina.az] Bazadan {max_items} ədəd elan axtarılır...")
        loc_id = None
        area_display = region_name

        if region_name in HIERARCHY_MAP:
            r_info = HIERARCHY_MAP[region_name]
            if sublocation_name and sublocation_name != "Bütün rayon üzrə" and sublocation_name in r_info.get("sublocations", {}):
                loc_id = r_info["sublocations"][sublocation_name]
                area_display = f"{region_name} ({sublocation_name})"
            else:
                loc_id = r_info.get("id")

        filter_obj: Dict[str, Any] = {
            'cityId': 1,
            'leased': (listing_type == "Kirayə")
        }
        if loc_id:
            filter_obj['locationId'] = loc_id

        cat_id = CATEGORIES_MAP.get(category_name, 1)
        if cat_id:
            filter_obj['categoryId'] = cat_id

        if room_ids and len(room_ids) > 0:
            filter_obj['roomIds'] = [str(r) for r in room_ids]

        if price_min is not None and float(price_min) > 0:
            filter_obj['priceFrom'] = int(float(price_min))
        if price_max is not None and float(price_max) > 0:
            filter_obj['priceTo'] = int(float(price_max))

        if area_min is not None and float(area_min) > 0:
            filter_obj['areaFrom'] = float(area_min)
        if area_max is not None and float(area_max) > 0:
            filter_obj['areaTo'] = float(area_max)

        if has_bill_of_sale:
            filter_obj['hasBillOfSale'] = True
        if has_mortgage:
            filter_obj['hasMortgage'] = True
        if has_repair:
            filter_obj['hasRepair'] = True

        cursor = None
        page_num = 1

        while (len(leads) - start_count) < max_items and not self.is_cancelled:
            conn = self.fetch_search_page(filter_obj, cursor=cursor, batch_size=24)
            if not conn:
                break

            edges = conn.get('edges', [])
            page_info = conn.get('pageInfo', {})
            has_next = page_info.get('hasNextPage', False)
            cursor = page_info.get('endCursor')

            if not edges:
                break

            for edge in edges:
                if self.is_cancelled or (len(leads) - start_count) >= max_items:
                    break

                node = edge.get('node', {})
                item_id = str(node.get('id', ''))
                self.stats["scanned"] += 1
                self.callback("stats", self.stats)

                is_business = node.get('isBusiness', False)
                company_info = node.get('company') or {}
                company_target_type = company_info.get('targetType')

                if target_type == "FSBO":
                    if is_business or company_target_type in ['AGENCY', 'RESIDENCE']:
                        self.stats["skipped"] += 1
                        self.callback("stats", self.stats)
                        continue

                self._sleep_jitter(0.6, 1.1)
                detail = self.fetch_item_detail(item_id)
                if not detail:
                    continue

                contact_type = (detail.get('contactTypeName') or '').strip().lower()
                contact_name = (detail.get('contactName') or '').strip()
                company_obj = detail.get('company') or {}
                business_obj = detail.get('business') or {}

                if target_type == "FSBO":
                    if contact_type == 'vasitəçi (agent)' or 'vasitəçi' in contact_type or 'agent' in contact_type:
                        self.stats["skipped"] += 1
                        self.callback("stats", self.stats)
                        continue
                    if company_obj.get('targetType') in ['AGENCY', 'RESIDENCE'] or business_obj:
                        self.stats["skipped"] += 1
                        self.callback("stats", self.stats)
                        continue
                    if self.is_agent_by_name(contact_name):
                        self.stats["skipped"] += 1
                        self.callback("stats", self.stats)
                        continue
                    if contact_type != 'mülkiyyətçi':
                        self.stats["skipped"] += 1
                        self.callback("stats", self.stats)
                        continue
                else: # AGENT mode
                    is_real_agent = (
                        contact_type == 'vasitəçi (agent)' or
                        'vasitəçi' in contact_type or
                        'agent' in contact_type or
                        company_obj.get('targetType') == 'AGENCY' or
                        (business_obj and business_obj.get('name'))
                    )
                    if not is_real_agent:
                        self.stats["skipped"] += 1
                        self.callback("stats", self.stats)
                        continue

                phone = self.fetch_unmasked_phone(item_id)
                if not phone or len(phone) < 7:
                    self.stats["skipped"] += 1
                    self.callback("stats", self.stats)
                    continue

                phone_clean = re.sub(r'[\s\(\)-]', '', phone)
                if phone_clean in self.phone_seen:
                    self.stats["duplicates"] += 1
                    self.callback("stats", self.stats)
                    continue

                self.phone_seen.add(phone_clean)

                price_val = (detail.get('price') or {}).get('total', 0)
                area_val = (detail.get('area') or {}).get('value', 0)
                rooms_val = detail.get('rooms')
                loc_obj = detail.get('location') or {}
                loc_name = loc_obj.get('fullName') or area_display

                # Price Trend & Discount detection
                desc_text = (detail.get('description') or '').lower()
                price_trend = "Bazar Qiyməti"
                if any(w in desc_text for w in ["endirim", "endirildi", "qiymət düşdü", "qiymeti dusdu", "ucuzlaşdı", "son qiymət"]):
                    price_trend = "🔻 Endirimli"
                elif any(w in desc_text for w in ["təcili", "tecili", "təcili satılır", "tecili satilir", "təcili pul lazımdır"]):
                    price_trend = "🔥 Təcili Satış"
                elif any(w in desc_text for w in ["hissə-hissə", "daxili kredit", "faizsiz"]):
                    price_trend = "💳 Daxili Kredit"

                sqm_val = round(price_val / float(area_val)) if area_val and float(area_val) > 0 and price_val > 0 else "-"

                agency_str = ""
                if target_type == "AGENT":
                    if business_obj and business_obj.get('name'):
                        agency_str = business_obj.get('name')
                    elif company_obj and company_obj.get('targetType') == 'AGENCY' and company_obj.get('name') != contact_name:
                        agency_str = company_obj.get('name')
                    else:
                        agency_str = "Fərdi Vasitəçi"

                lead_record = {
                    "№": len(leads) + 1,
                    "Mənbə": "bina.az",
                    "Ad / Mülkiyyətçi": contact_name or ("Ev Sahibi" if target_type == "FSBO" else "Vasitəçi"),
                    "Əlaqə Nömrəsi": phone,
                    "Status": "Mülkiyyətçi" if target_type == "FSBO" else "Vasitəçi (agent)",
                    "Agentlik / Şirkət": agency_str if target_type == "AGENT" else "-",
                    "Qiymət (AZN)": price_val,
                    "m² Qiyməti (AZN)": sqm_val,
                    "Qiymət Trendi": price_trend,
                    "Otaq": rooms_val,
                    "Sahə (m²)": area_val,
                    "Yerləşmə": loc_name,
                    "Kupça": "Var" if detail.get('hasBillOfSale') else "Yoxdur",
                    "İpoteka": "Var" if detail.get('hasMortgage') else "Yoxdur",
                    "Təmir": "Var" if detail.get('hasRepair') else "Yoxdur",
                    "Elan ID": item_id,
                    "Elan Keçidi": f"https://bina.az/items/{item_id}",
                    "Tarix": detail.get('updatedAt', '')[:10]
                }

                leads.append(lead_record)
                self.stats["accepted"] = len(leads)
                self.callback("stats", self.stats)
                self.callback("lead", lead_record)

                self.callback("log", f"  ✅ [bina.az #{len(leads) - start_count}/{max_items}] {lead_record['Ad / Mülkiyyətçi']} | {phone} | {price_val} AZN | {loc_name}")

            if not has_next:
                break
            page_num += 1

        return len(leads) - start_count

    def run(self,
            target_type: str = "FSBO",
            region_name: str = "Bütün Bakı",
            sublocation_name: Optional[str] = None,
            category_name: str = "Bütün Mənzillər",
            listing_type: str = "Alqı-satqı",
            room_ids: Optional[List[str]] = None,
            price_min: Optional[int] = None,
            price_max: Optional[int] = None,
            area_min: Optional[float] = None,
            area_max: Optional[float] = None,
            has_bill_of_sale: Optional[bool] = None,
            has_mortgage: Optional[bool] = None,
            has_repair: Optional[bool] = None,
            bina_target: int = 20,
            yeni_target: int = 10,
            tap_target: int = 0,
            output_dir: str = ".") -> Dict[str, Any]:

        self.is_cancelled = False
        self.stats = {"scanned": 0, "accepted": 0, "skipped": 0, "duplicates": 0}
        self.phone_seen = set()
        leads: List[Dict[str, Any]] = []

        total_target = bina_target + yeni_target + tap_target
        if total_target <= 0:
            total_target = 30
            bina_target = 30

        mode_text = "Birbaşa Mülkiyyətçi (FSBO)" if target_type == "FSBO" else "Əmlak Agenti (Rieltor)"
        self.callback("log", f"🎯 Axtarış Rejimi: {mode_text}")
        self.callback("log", f"📍 Seçilmiş Ərazi: {region_name} | Növ: {listing_type}")
        self.callback("log", f"🎯 Hədəflər: bina.az: {bina_target} | yeniemlak.az: {yeni_target} | tap.az: {tap_target}")

        # 1. Scrape from bina.az
        if bina_target > 0 and not self.is_cancelled:
            self.scrape_bina(
                target_type=target_type,
                max_items=bina_target,
                leads=leads,
                region_name=region_name,
                sublocation_name=sublocation_name,
                category_name=category_name,
                listing_type=listing_type,
                room_ids=room_ids,
                price_min=price_min,
                price_max=price_max,
                area_min=area_min,
                area_max=area_max,
                has_bill_of_sale=has_bill_of_sale,
                has_mortgage=has_mortgage,
                has_repair=has_repair
            )

        # 2. Scrape from yeniemlak.az
        if yeni_target > 0 and not self.is_cancelled:
            self.scrape_yeniemlak(
                target_type=target_type,
                max_items=yeni_target,
                leads=leads,
                region_name=region_name,
                sublocation_name=sublocation_name,
                listing_type=listing_type
            )

        # 3. Scrape from tap.az
        if tap_target > 0 and not self.is_cancelled:
            self.scrape_tapaz(
                target_type=target_type,
                max_items=tap_target,
                leads=leads,
                region_name=region_name,
                sublocation_name=sublocation_name,
                category_name=category_name,
                listing_type=listing_type
            )

        # File Export
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        prefix = "bina_fsbo" if target_type == "FSBO" else "bina_agents"
        sanitized_area = re.sub(r'[^a-zA-Z0-9_]', '', area_display.replace(' ', '_'))
        base_name = f"{prefix}_{sanitized_area}_{timestamp}"
        
        excel_path = os.path.join(output_dir, f"{base_name}.xlsx")
        csv_path = os.path.join(output_dir, f"{base_name}.csv")
        json_path = os.path.join(output_dir, f"{base_name}.json")

        self.export_excel(leads, excel_path, target_type)
        self.export_csv(leads, csv_path)
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(leads, f, indent=2, ensure_ascii=False)

        self.callback("log", f"\n🎉 Tamamlandı! Cəmi {len(leads)} ədəd qeyd uğurla bazaya yığıldı.")
        self.callback("log", f"📁 Excel faylı: {os.path.abspath(excel_path)}")
        self.callback("finished", {
            "excel_path": os.path.abspath(excel_path),
            "csv_path": os.path.abspath(csv_path),
            "json_path": os.path.abspath(json_path),
            "count": len(leads)
        })

        return {
            "leads": leads,
            "excel_path": os.path.abspath(excel_path),
            "count": len(leads)
        }

    def export_excel(self, leads: List[Dict[str, Any]], excel_path: str, target_type: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mülkiyyətçilər" if target_type == "FSBO" else "Agentlər"
        ws.views.sheetView[0].showGridLines = True

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        
        row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        row_white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        regular_font = Font(name="Segoe UI", size=10, color="1E293B")
        bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
        name_font = Font(name="Segoe UI", size=10, bold=True, color="1E3A8A")
        phone_font = Font(name="Segoe UI", size=10, bold=True, color="15803D")
        link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")

        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )

        columns = [
            ("№", "№", 6, "center"),
            ("Mənbə", "Mənbə", 14, "center"),
            ("Ad / Mülkiyyətçi", "Ad / Mülkiyyətçi", 22, "left"),
            ("Əlaqə Nömrəsi", "Əlaqə Nömrəsi", 18, "center"),
            ("Status", "Status", 16, "center"),
            ("Agentlik / Şirkət", "Agentlik / Şirkət", 24, "left"),
            ("Qiymət (AZN)", "Qiymət (AZN)", 16, "right"),
            ("Otaq", "Otaq", 8, "center"),
            ("Sahə (m²)", "Sahə (m²)", 12, "center"),
            ("Yerləşmə", "Yerləşmə", 28, "left"),
            ("Kupça", "Kupça", 10, "center"),
            ("İpoteka", "İpoteka", 10, "center"),
            ("Təmir", "Təmir", 10, "center"),
            ("Elan ID", "Elan ID", 12, "center"),
            ("Elan Keçidi (Klikləyin)", "Elan Keçidi", 30, "left")
        ]

        ws.row_dimensions[1].height = 30
        for col_idx, (col_title, key, width, align) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_title)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        for row_idx, item in enumerate(leads, 2):
            ws.row_dimensions[row_idx].height = 24
            fill = row_alt_fill if row_idx % 2 == 0 else row_white_fill

            for col_idx, (col_title, key, width, align) in enumerate(columns, 1):
                val = item.get(key)
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = fill
                cell.border = thin_border
                cell.font = regular_font

                if key == "№":
                    cell.value = int(val)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif key == "Ad / Mülkiyyətçi":
                    cell.value = str(val or "")
                    cell.font = name_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif key == "Əlaqə Nömrəsi":
                    cell.value = str(val or "")
                    cell.font = phone_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif key == "Qiymət (AZN)":
                    try:
                        cell.value = float(val) if val else 0
                        cell.number_format = '#,##0 AZN'
                        cell.font = bold_font
                    except Exception:
                        cell.value = str(val or "")
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif key == "Sahə (m²)":
                    try:
                        cell.value = float(val) if val and str(val) != '-' else 0
                        cell.number_format = '#,##0.0'
                    except Exception:
                        cell.value = str(val or "")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif key == "Elan Keçidi":
                    url_str = str(val or "")
                    cell.value = url_str
                    cell.hyperlink = url_str
                    cell.font = link_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif key in ["Kupça", "İpoteka", "Təmir"]:
                    cell.value = str(val or "")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if str(val) == "Var":
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="166534")
                    else:
                        cell.font = Font(name="Segoe UI", size=10, color="94A3B8")
                else:
                    cell.value = str(val or "")
                    cell.alignment = Alignment(horizontal=align, vertical="center")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(excel_path)

    def export_csv(self, leads: List[Dict[str, Any]], csv_path: str):
        import pandas as pd
        df = pd.DataFrame(leads)
        df.to_csv(csv_path, index=False, encoding='utf-8-sig')
